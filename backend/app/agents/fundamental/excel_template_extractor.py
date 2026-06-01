"""
Deterministic extractor for BUMI-style template financial Excel files.

Template format:
  Row N:   PARAMETER | YEAR1 | YEAR2 | YEAR3 | Satuan
  ...
  Data rows: label in col 0, numeric values in year cols

Supports any number of year columns. No LLM required.
"""
import io
import logging
import re
from typing import Any, Dict, List, Optional, Tuple

from app.agents.fundamental.state import FinancialYear

logger = logging.getLogger(__name__)


def _normalize(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip()).upper()


# ── Field map: (normalized_pattern, section, field) ───────────────────────────
# More specific patterns listed FIRST to prevent shorter patterns from winning.
_FIELD_MAP: List[Tuple[str, str, str]] = [
    # Cash-flow ending/beginning cash — must precede generic "KAS DAN SETARA KAS"
    ("KAS DAN SETARA KAS PADA AKHIR",       "cash_flow",        "ending_cash"),
    ("KAS DAN SETARA KAS PADA AWAL",        "cash_flow",        "beginning_cash"),

    # Income Statement
    ("LABA BRUTO",                          "income_statement", "gross_profit"),
    ("LABA USAHA",                          "income_statement", "operating_income"),
    ("LABA SEBELUM PAJAK",                  "income_statement", "pretax_income"),
    ("LABA SETELAH PAJAK",                  "income_statement", "net_income_after_tax"),
    ("LABA TAHUN BERJALAN - NETO",          "income_statement", "net_income"),
    ("LABA TAHUN BERJALAN - PEMILIK",       "income_statement", "parent_net_income"),
    ("LABA PER 1.000 SAHAM",               "income_statement", "eps"),
    ("LABA PER SAHAM",                      "income_statement", "eps"),
    ("JUMLAH LABA KOMPREHENSIF",            "income_statement", "comprehensive_income"),
    ("BEBAN POKOK PENDAPATAN",              "income_statement", "cogs"),
    ("BEBAN USAHA",                         "income_statement", "operating_expenses"),
    ("BEBAN PAJAK PENGHASILAN",             "income_statement", "tax_expense"),
    ("BEBAN BUNGA DAN KEUANGAN",            "income_statement", "interest_expense"),
    ("PENGHASILAN BUNGA",                   "income_statement", "interest_income"),
    ("BAGIAN LABA ENTITAS ASOSIASI",        "income_statement", "associates_earnings"),
    ("LABA (RUGI) SELISIH KURS",            "income_statement", "forex_gain_loss"),
    ("PENDAPATAN",                          "income_statement", "revenue"),

    # Balance Sheet — Assets
    ("KAS DAN SETARA KAS",                  "balance_sheet",    "cash"),
    ("KAS DI BANK DIBATASI",                "balance_sheet",    "restricted_cash"),
    ("PIUTANG USAHA - PIHAK KETIGA",        "balance_sheet",    "accounts_receivable"),
    ("PIUTANG USAHA - PIHAK BERELASI",      "balance_sheet",    "related_receivables"),
    ("PERSEDIAAN",                          "balance_sheet",    "inventory"),
    ("JUMLAH ASET TIDAK LANCAR",            "balance_sheet",    "noncurrent_assets"),
    ("JUMLAH ASET LANCAR",                  "balance_sheet",    "current_assets"),
    ("JUMLAH ASET",                         "balance_sheet",    "total_assets"),
    # "PEMBELIAN ASET TETAP" must precede "ASET TETAP" to prevent substring collision
    ("PEMBELIAN ASET TETAP",                "cash_flow",        "capex"),
    ("ASET TETAP",                          "balance_sheet",    "fixed_assets"),
    ("PROPERTI PERTAMBANGAN",               "balance_sheet",    "mining_properties"),
    ("GOODWILL",                            "balance_sheet",    "goodwill"),

    # Balance Sheet — Liabilities (specific before generic)
    ("PINJAMAN JANGKA PANJANG (JATUH TEMPO","balance_sheet",    "current_ltdebt"),
    ("UTANG OBLIGASI (JATUH TEMPO",         "balance_sheet",    "current_bonds"),
    ("PINJAMAN JANGKA PANJANG (NET",        "balance_sheet",    "long_term_loans"),
    ("UTANG OBLIGASI (JANGKA PANJANG",      "balance_sheet",    "long_term_bonds"),
    ("PINJAMAN JANGKA PENDEK",              "balance_sheet",    "short_term_loans"),
    ("UTANG USAHA - PIHAK KETIGA",          "balance_sheet",    "trade_payables"),
    ("BEBAN AKRUAL",                        "balance_sheet",    "accrued_expenses"),
    ("UTANG PAJAK",                         "balance_sheet",    "tax_payables"),
    ("JUMLAH LIABILITAS JANGKA PENDEK",     "balance_sheet",    "current_liabilities"),
    ("JUMLAH LIABILITAS JANGKA PANJANG",    "balance_sheet",    "noncurrent_liabilities"),
    ("JUMLAH LIABILITAS DAN EKUITAS",       "balance_sheet",    "total_liabilities_and_equity"),
    ("JUMLAH LIABILITAS",                   "balance_sheet",    "total_liabilities"),

    # Balance Sheet — Equity
    ("JUMLAH EKUITAS - PEMILIK",            "balance_sheet",    "parent_equity"),
    ("KEPENTINGAN NONPENGENDALI",           "balance_sheet",    "nci"),
    ("EKUITAS - NETO",                      "balance_sheet",    "total_equity"),

    # Cash Flow
    ("ARUS KAS NETO DARI (UNTUK) AKTIVITAS OPERASI",   "cash_flow", "operating_cash_flow"),
    ("ARUS KAS NETO DARI (UNTUK) AKTIVITAS INVESTASI", "cash_flow", "net_investing_cf"),
    ("ARUS KAS NETO DARI AKTIVITAS PENDANAAN",         "cash_flow", "net_financing_cf"),
    ("PENERIMAAN DARI PELANGGAN",           "cash_flow",        "customer_receipts"),
    ("KENAIKAN (PENURUNAN) NETO KAS",       "cash_flow",        "net_cash_change"),
]

# These fields are stored as negative in the template → convert to absolute value
_ABS_FIELDS = {
    "cogs", "operating_expenses", "tax_expense",
    "interest_expense", "capex",
}


# ── Public API ────────────────────────────────────────────────────────────────

def is_template_format(content: bytes) -> bool:
    """Return True if the xlsx matches the PARAMETER/year header structure."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(
            io.BytesIO(content), data_only=True, read_only=True
        )
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
                if not row or row[0] is None:
                    continue
                if _normalize(row[0]) == "PARAMETER":
                    if any(_to_year(c) is not None for c in row[1:] if c is not None):
                        return True
        return False
    except Exception:
        return False


def extract_from_template(content: bytes) -> List[FinancialYear]:
    """
    Deterministically extract FinancialYear list from a BUMI-style template.
    Returns an empty list if parsing fails.
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sheet_name in wb.sheetnames:
            result = _parse_sheet(wb[sheet_name])
            if result:
                logger.info(
                    "Template extractor: %d year(s) from sheet '%s'",
                    len(result), sheet_name,
                )
                return result
    except Exception as exc:
        logger.error("Template extractor error: %s", exc)
    return []


# ── Internal helpers ──────────────────────────────────────────────────────────

def _to_year(v: Any) -> Optional[int]:
    try:
        y = int(str(v).strip())
        return y if 2000 <= y <= 2100 else None
    except (ValueError, TypeError):
        return None


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.startswith("=") or s == "":
        return None
    try:
        return float(s.replace(",", "").replace(" ", ""))
    except (ValueError, TypeError):
        return None


def _match_field(label_upper: str) -> Optional[Tuple[str, str]]:
    for pattern, section, field in _FIELD_MAP:
        if pattern in label_upper:
            return section, field
    return None


def _parse_sheet(ws) -> List[FinancialYear]:
    rows = list(ws.iter_rows(values_only=True))

    # Find the PARAMETER header row and year columns
    header_idx: Optional[int] = None
    year_cols: List[int] = []
    year_vals: List[int] = []

    for i, row in enumerate(rows):
        if not row or row[0] is None:
            continue
        if _normalize(row[0]) == "PARAMETER":
            for j, cell in enumerate(row[1:], start=1):
                y = _to_year(cell)
                if y is not None:
                    year_cols.append(j)
                    year_vals.append(y)
            if year_vals:
                header_idx = i
                break

    if header_idx is None or not year_vals:
        return []

    # Per-year data buckets
    data: List[Dict[str, Dict[str, Optional[float]]]] = [
        {"income_statement": {}, "balance_sheet": {}, "cash_flow": {}}
        for _ in year_vals
    ]

    for row in rows[header_idx + 1:]:
        if not row or row[0] is None:
            continue
        label_upper = _normalize(row[0])
        if not label_upper:
            continue
        match = _match_field(label_upper)
        if match is None:
            continue
        section, field = match

        for i, col_idx in enumerate(year_cols):
            if col_idx >= len(row):
                continue
            val = _to_float(row[col_idx])
            if val is None:
                continue
            if field in _ABS_FIELDS:
                val = abs(val)
            # Only write if not already set (first match wins due to priority order)
            if field not in data[i][section]:
                data[i][section][field] = val

    results: List[FinancialYear] = []
    for i, year in enumerate(year_vals):
        bs = data[i]["balance_sheet"]
        is_ = data[i]["income_statement"]
        cf = data[i]["cash_flow"]

        # Derive total_debt from components
        debt_parts = [
            bs.get("short_term_loans"),
            bs.get("current_ltdebt"),
            bs.get("current_bonds"),
            bs.get("long_term_loans"),
            bs.get("long_term_bonds"),
        ]
        clean_debt = [v for v in debt_parts if v is not None]
        if clean_debt:
            bs["total_debt"] = sum(clean_debt)

        # Derive free_cash_flow
        ocf = cf.get("operating_cash_flow")
        capex = cf.get("capex")
        if ocf is not None and capex is not None:
            cf["free_cash_flow"] = ocf - capex

        results.append(
            FinancialYear(year=year, income_statement=is_, balance_sheet=bs, cash_flow=cf)
        )

    return results
